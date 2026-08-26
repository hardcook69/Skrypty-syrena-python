#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Syrena — Wprowadzanie usług pracowników z Excela (GUI)
Domyślnie config_testowy.json obok skryptu; inny plik jako argument CLI
(np. `python dodaj_uslugi_pracownikow_gui.py config_produkcja.json` dla
produkcji). Ten sam mechanizm co dodaj_uslugi_pracownikow.py (wersja CLI)
— patrz tam po szczegóły przepływu API i kształtu payloadu
POST :5000/api/dictionary-value.

Na razie obsługuje TYLKO usługi pracowników (Typ usługi = "Dla pracownika").

BEZPIECZEŃSTWO: to skrypt PISZĄCY do systemu. "Pokaż podgląd" nic nie
wysyła. Wysyłka wymaga jawnego potwierdzenia w oknie dialogowym z
liczbą usług (i dodatkowym ostrzeżeniem, gdy środowisko to PRODUKCJA).
Usługi, które już istnieją (taka sama treść w słowniku 51) są pomijane
— bezpiecznie uruchomić kilka razy na tym samym pliku.
"""

import tkinter as tk
from tkinter import messagebox, filedialog
import threading, requests, json, os, sys, traceback, logging

try:
    import ttkbootstrap as ttkb
    HAS_TTKB = True
except ImportError:
    HAS_TTKB = False

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


# ── Konfiguracja ─────────────────────────────────────────────────────────────
SCRIPT_DIR      = os.path.dirname(os.path.abspath(__file__))
CONFIG_NAME     = sys.argv[1] if len(sys.argv) > 1 else "config_testowy.json"
CONFIG_PATH     = os.path.join(SCRIPT_DIR, CONFIG_NAME)
CRASH_LOG_PATH  = os.path.join(SCRIPT_DIR, "dodaj_uslugi_pracownikow_gui_crash.log")
LAST_LOGIN_PATH = os.path.join(SCRIPT_DIR, "dodaj_uslugi_pracownikow_gui_last_login.json")


def _fatal_startup_error(title, message, exc=None):
    detail = f"{message}\n\n{traceback.format_exc()}" if exc else message
    try:
        with open(CRASH_LOG_PATH, "a", encoding="utf-8") as f:
            from datetime import datetime as _dt
            f.write(f"\n[{_dt.now().strftime('%Y-%m-%d %H:%M:%S')}] {title}\n{detail}\n")
    except Exception:
        pass
    try:
        root = tk.Tk()
        root.withdraw()
        messagebox.showerror(title, f"{message}\n\nSzczegóły zapisane w:\n{CRASH_LOG_PATH}")
        root.destroy()
    except Exception:
        print(f"{title}: {detail}", file=sys.stderr)
    sys.exit(1)


if not HAS_TTKB:
    _fatal_startup_error(
        "Brak biblioteki ttkbootstrap",
        "To GUI wymaga pakietu 'ttkbootstrap'.\nZainstaluj go poleceniem:\n\n"
        "    pip install ttkbootstrap")

if not HAS_OPENPYXL:
    _fatal_startup_error(
        "Brak biblioteki openpyxl",
        "To GUI wymaga pakietu 'openpyxl' (odczyt plików Excel).\nZainstaluj go poleceniem:\n\n"
        "    pip install openpyxl")


try:
    with open(CONFIG_PATH, encoding="utf-8") as _f:
        CFG = json.load(_f)
except FileNotFoundError:
    _fatal_startup_error(
        "Brak pliku konfiguracyjnego",
        f"Nie znaleziono {CONFIG_PATH}.\nUpewnij się, że {CONFIG_NAME} "
        f"leży w tym samym folderze co ten skrypt.")
except json.JSONDecodeError as e:
    _fatal_startup_error(
        f"Błędny format {CONFIG_NAME}",
        f"Plik {CONFIG_PATH} nie jest poprawnym JSON-em: {e}", exc=e)

try:
    LOG_PATH = os.path.join(SCRIPT_DIR, CFG.get("log_file", "dodaj_uslugi_pracownikow_gui.log"))
    logging.basicConfig(filename=LOG_PATH, level=logging.DEBUG,
                        format="%(asctime)s %(levelname)s %(message)s",
                        encoding="utf-8")

    AUTH_URL       = CFG["auth_url"]
    EMP_URL        = CFG.get("employee_api_url", AUTH_URL.replace(":5010", ":5000"))
    ORG_ID_DEFAULT = CFG.get("organization_id", 1)
    ENV_LABEL      = CFG.get("env_label", "TEST")
except KeyError as e:
    _fatal_startup_error(
        f"Brak klucza w {CONFIG_NAME}",
        f"W pliku konfiguracyjnym brakuje wymaganego klucza: {e}\nWymagane: auth_url.", exc=e)


def _load_last_login():
    try:
        with open(LAST_LOGIN_PATH, encoding="utf-8") as f:
            return (json.load(f).get("last_login") or "").strip()
    except Exception:
        return ""

def _save_last_login(login_value):
    try:
        with open(LAST_LOGIN_PATH, "w", encoding="utf-8") as f:
            json.dump({"last_login": login_value}, f)
    except Exception as e:
        logging.warning(f"Nie udało się zapisać ostatniego loginu: {e}")


# ── Styl (ttkbootstrap) ───────────────────────────────────────────────────────
IS_PROD = ENV_LABEL.strip().upper() in ("PROD", "PRODUKCJA", "PRODUCTION")
if IS_PROD:
    ACCENT_STYLE = "danger"
elif ENV_LABEL == "TEST":
    ACCENT_STYLE = "primary"
else:
    ACCENT_STYLE = "info"
THEME_NAME = "united"

FONT_TITLE   = ("Segoe UI", 12, "bold")
FONT_SECTION = ("Segoe UI", 10, "bold")
FONT_BASE    = ("Segoe UI", 10)
FONT_SMALL   = ("Segoe UI", 9)
FONT_TINY    = ("Segoe UI", 8)


def _card(parent):
    return ttkb.Labelframe(parent, text="", bootstyle="secondary")


# ── Backend (ten sam mechanizm co dodaj_uslugi_pracownikow.py) ───────────────
TIMEOUT = 20
USLUGI_DICTIONARY_ID = 51
DEFAULT_SHEET = "Arkusz1"
DEFAULT_TYP_USLUGI = "Dla pracownika"
DEFAULT_POCHODZENIE = "Wewnętrzna"

COL_USLUGA_HEADER = "Usługa pracownik"
COL_KATEGORIA_PREFIX = "Kategoria usługi("
COL_KATEGORIA_SUFFIX = "(pracownik)"


class Client:
    def __init__(self):
        self.servers = {"auth": AUTH_URL, "employee": EMP_URL}
        self.s = requests.Session()
        self.s.headers.update({"Content-Type": "application/json", "Accept": "application/json"})
        self._creds = None
        self._current_org_id = None

    def login(self, username, password):
        for key in ("login", "userName", "username"):
            try:
                r = self.s.post(f"{self.servers['auth']}/api/authentication/token",
                                 json={key: username, "password": password}, timeout=TIMEOUT)
                if r.status_code == 200:
                    d = r.json()
                    tok = d.get("token") or d.get("access_token")
                    if tok:
                        self.s.headers["Authorization"] = f"Bearer {tok}"
                        self._creds = (username, password)
                        logging.info(f"Zalogowano: {d.get('userName')}")
                        return True
            except Exception as e:
                logging.error(f"Błąd logowania (klucz={key}): {e}")
        return False

    def select_organization(self, org_id):
        try:
            r = self.s.post(f"{self.servers['auth']}/api/authentication/token",
                             json={"organizationId": org_id}, timeout=TIMEOUT)
        except Exception as e:
            logging.error(f"Błąd przełączania organizacji org_id={org_id}: {e}")
            return False
        if r.status_code != 200:
            logging.error(f"Przełączenie na organizationId={org_id} nieudane: HTTP {r.status_code}")
            return False
        try:
            d = r.json()
        except Exception:
            d = {}
        tok = d.get("token") or d.get("access_token")
        if tok:
            self.s.headers["Authorization"] = f"Bearer {tok}"
        self._current_org_id = org_id
        logging.info(f"Przełączono na organizationId={org_id}")
        return True

    def _relogin(self):
        if not self._creds:
            return False
        if self.login(*self._creds):
            if self._current_org_id is not None:
                self.select_organization(self._current_org_id)
            return True
        return False

    def _get(self, url, params=None, _retry=True):
        logging.debug(f"GET {url} params={params}")
        try:
            r = self.s.get(url, params=params, timeout=TIMEOUT)
        except Exception as e:
            logging.error(f"Wyjątek przy GET {url}: {e}")
            return None, str(e)
        logging.debug(f"-> HTTP {r.status_code} {url}: {r.text[:500]}")
        if r.status_code == 401 and _retry and self._relogin():
            return self._get(url, params=params, _retry=False)
        return r, None

    def _post(self, url, json_body=None, _retry=True):
        logging.debug(f"POST {url} body={json_body}")
        try:
            r = self.s.post(url, json=json_body, timeout=TIMEOUT)
        except Exception as e:
            logging.error(f"Wyjątek przy POST {url}: {e}")
            return None, str(e)
        logging.debug(f"-> HTTP {r.status_code} {url}: {r.text[:500]}")
        if r.status_code == 401 and _retry and self._relogin():
            return self._post(url, json_body=json_body, _retry=False)
        return r, None


def fetch_structure_elements(client, dictionary_id):
    r, err = client._get(f"{client.servers['employee']}/api/dictionary-structure-element/by-dictionary-id/{dictionary_id}")
    if err or r is None or r.status_code != 200:
        logging.error(f"Błąd pobierania struktury słownika {dictionary_id}: {err or (r and r.status_code)}")
        return []
    data = r.json()
    items = data if isinstance(data, list) else (data.get("items") or data.get("data") or [])
    logging.info(f"Słownik {dictionary_id}: {len(items)} pól struktury")
    return items


def fetch_dictionary_values(client, dictionary_id):
    r, err = client._get(f"{client.servers['employee']}/api/dictionary-value/by-dictionary-id/{dictionary_id}")
    if err or r is None or r.status_code != 200:
        logging.error(f"Błąd pobierania wartości słownika {dictionary_id}: {err or (r and r.status_code)}")
        return []
    data = r.json()
    return data if isinstance(data, list) else (data.get("items") or data.get("data") or [])


def find_value_id_by_content(values, wanted_content, field_label):
    """GUI-owa wersja: zamiast przerywać program, zwraca (id, błąd)."""
    wanted_norm = wanted_content.strip().lower()
    for v in values:
        if (v.get("content") or "").strip().lower() == wanted_norm:
            return v.get("id"), None
    available = ", ".join(repr(v.get("content")) for v in values)
    return None, f"Nie znaleziono wartości '{wanted_content}' dla pola '{field_label}'. Dostępne opcje: {available}"


def resolve_value_attributes(client, structure_elements, kategoria_uslugi_content,
                              typ_uslugi_content, pochodzenie_content, dict_value_cache=None):
    """Buduje valueAttributes na podstawie struktury słownika 51. Zwraca
    (attrs, błąd) — pole, którego nazwa nie pasuje do znanej reguły
    (Typ/Kategoria/Pochodzenie usługi), zwraca błąd zamiast zgadywać."""
    attrs = []
    if dict_value_cache is None:
        dict_value_cache = {}
    for elem in structure_elements:
        name = (elem.get("name") or "").strip()
        linked_dict_id = elem.get("elementDictionaryId")
        attr_kind = elem.get("elementKind")
        attr_type = elem.get("elementType")

        if name.lower() == "typ usługi":
            wanted = typ_uslugi_content
        elif name.lower() == "kategoria usługi":
            wanted = kategoria_uslugi_content
        elif name.lower() == "pochodzenie usługi":
            wanted = pochodzenie_content
        else:
            return None, (f"Słownik {USLUGI_DICTIONARY_ID} wymaga pola '{name}', którego to GUI "
                          f"jeszcze nie obsługuje (obsługiwane: Typ usługi, Kategoria usługi, "
                          f"Pochodzenie usługi).")

        if linked_dict_id not in dict_value_cache:
            dict_value_cache[linked_dict_id] = fetch_dictionary_values(client, linked_dict_id)
        value_id, err = find_value_id_by_content(dict_value_cache[linked_dict_id], wanted, name)
        if err:
            return None, err

        attrs.append({
            "id": 0, "rowVersion": 0, "isDeleted": False,
            "attributeKind": attr_kind, "attributeType": attr_type, "value": value_id,
        })
    return attrs, None


def post_service(client, content, display_order, value_attributes):
    payload = {
        "id": 0, "rowVersion": 0, "isDeleted": False,
        "dictionaryId": USLUGI_DICTIONARY_ID,
        "content": content, "displayOrder": display_order,
        "parentDictionaryValueId": None,
        "valueAttributes": value_attributes,
    }
    r, err = client._post(f"{client.servers['employee']}/api/dictionary-value", json_body=payload)
    if err:
        return False, err
    if r.status_code not in (200, 201, 204):
        try:
            detail = r.text[:300]
        except Exception:
            detail = ""
        return False, f"HTTP {r.status_code}: {detail}"
    return True, None


def list_sheet_names(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    return wb.sheetnames


def _build_merge_map(ws):
    m = {}
    for rng in ws.merged_cells.ranges:
        top_val = ws.cell(row=rng.min_row, column=rng.min_col).value
        for row in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                m[(row, col)] = top_val
    return m


def _cell_value(ws, merge_map, row, col):
    v = ws.cell(row=row, column=col).value
    if v is None:
        v = merge_map.get((row, col))
    return v


def find_header_column(ws, header_row, predicate):
    for col in range(1, ws.max_column + 1):
        header = ws.cell(row=header_row, column=col).value
        if header and predicate(str(header)):
            return col
    return None


def read_employee_services_from_excel(xlsx_path, sheet_name):
    """GUI-owa wersja: zwraca (lista_usług, błąd) zamiast przerywać program."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        return None, f"Arkusz '{sheet_name}' nie istnieje w pliku. Dostępne: {wb.sheetnames}"
    ws = wb[sheet_name]
    merge_map = _build_merge_map(ws)

    col_usluga = find_header_column(ws, 1, lambda h: h.strip() == COL_USLUGA_HEADER)
    if col_usluga is None:
        return None, f"Nie znaleziono kolumny '{COL_USLUGA_HEADER}' w wierszu nagłówka arkusza '{ws.title}'."
    col_kategoria = find_header_column(
        ws, 1, lambda h: h.startswith(COL_KATEGORIA_PREFIX) and h.rstrip().endswith(COL_KATEGORIA_SUFFIX))
    if col_kategoria is None:
        return None, f"Nie znaleziono kolumny kategorii usługi (pracownik) w wierszu nagłówka arkusza '{ws.title}'."

    seen = {}
    order = []
    for row in range(2, ws.max_row + 1):
        usluga = _cell_value(ws, merge_map, row, col_usluga)
        kategoria = _cell_value(ws, merge_map, row, col_kategoria)
        if not usluga or not str(usluga).strip():
            continue
        usluga = str(usluga).strip()
        kategoria = str(kategoria).strip() if kategoria else ""
        if not kategoria:
            return None, f"Wiersz {row}: usługa '{usluga}' nie ma podanej kategorii usługi."
        norm = usluga.lower()
        if norm in seen:
            prev_content, prev_kategoria, prev_row = seen[norm]
            if prev_kategoria.lower() != kategoria.lower():
                return None, (f"Usługa '{usluga}' ma różne kategorie w wierszach {prev_row} "
                              f"('{prev_kategoria}') i {row} ('{kategoria}') — popraw Excel.")
            continue
        seen[norm] = (usluga, kategoria, row)
        order.append(norm)

    return [seen[n] for n in order], None


# ─────────────────────────────────────────────────────────────────────────────
class App(ttkb.Window):
    def __init__(self):
        super().__init__(title=f"Syrena — Usługi pracowników z Excela [{ENV_LABEL}]",
                          themename=THEME_NAME, size=(950, 760), resizable=(True, True))

        self.client = Client()
        self._xlsx_path = None
        self._pending_services = []   # [(content, kategoria, row), ...] do utworzenia
        self._skipped_existing = []
        self._structure_elements = []

        self._build_ui()

    # ── BUILD UI ──────────────────────────────────────────────────────────────
    def _build_ui(self):
        env_icon = "⚠ " if IS_PROD else ""
        hdr = ttkb.Frame(self, bootstyle=f"@{ACCENT_STYLE}", height=52)
        hdr.pack(fill="x"); hdr.pack_propagate(False)
        ttkb.Label(hdr, text="  Syrena — Usługi pracowników z Excela", bootstyle=f"@{ACCENT_STYLE}",
                   font=FONT_TITLE).pack(side="left", padx=(12, 4))
        ttkb.Label(hdr, text=f"{env_icon}[{ENV_LABEL}]", bootstyle=f"@{ACCENT_STYLE}",
                   font=("Segoe UI", 12, "bold")).pack(side="left")
        self._lbl_status = ttkb.Label(hdr, text="●  Niezalogowany", bootstyle=f"@{ACCENT_STYLE}", font=FONT_SMALL)
        self._lbl_status.pack(side="right", padx=12)
        if IS_PROD:
            ttkb.Label(hdr, text="PRODUKCJA — usługi trafią od razu do systemu",
                       bootstyle=f"@{ACCENT_STYLE}", font=FONT_TINY).pack(side="right", padx=12)

        nb = ttkb.Notebook(self, bootstyle=ACCENT_STYLE)
        nb.pack(fill="both", expand=True, padx=10, pady=8)
        self._nb = nb

        t1 = ttkb.Frame(nb); nb.add(t1, text="  1. Logowanie  ")
        t2 = ttkb.Frame(nb); nb.add(t2, text="  2. Import usług  ")

        self._tab_login(t1)
        self._tab_import(t2)

        lf = ttkb.Labelframe(self, text=" Log ", bootstyle="secondary")
        lf.pack(fill="x", padx=10, pady=(0, 8))
        self._log = ttkb.ScrolledText(
            lf, height=8, font=("Consolas", 9), bg="#1e1e1e", fg="#ccc",
            insertbackground="white", state="disabled", wrap="word", auto_hide=True)
        self._log.pack(fill="x", padx=6, pady=4)
        self._log.tag_config("ok", foreground="#6dbf67")
        self._log.tag_config("err", foreground="#e06c6c")
        self._log.tag_config("inf", foreground="#9cdcfe")

    # ── TAB 1: Logowanie ──────────────────────────────────────────────────────
    def _tab_login(self, p):
        outer = ttkb.Frame(p, padding=(24, 20))
        outer.pack(anchor="nw", fill="x")
        ttkb.Label(outer, text="Zaloguj się do API Syrena",
                   font=FONT_TITLE).pack(anchor="w", pady=(0, 4))
        ttkb.Label(outer, text=f"Środowisko: {ENV_LABEL}   •   Host: {AUTH_URL}",
                   bootstyle="secondary", font=FONT_TINY).pack(anchor="w", pady=(0, 14))

        card = _card(outer)
        card.pack(anchor="w")
        f = ttkb.Frame(card, padding=(20, 18))
        f.pack()

        for row, (lbl, attr, kw) in enumerate([
            ("Login:", "_e_user", {}),
            ("Hasło:", "_e_pass", {"show": "*"}),
            ("Organizacja (ID):", "_e_org", {}),
        ]):
            ttkb.Label(f, text=lbl, font=FONT_BASE).grid(row=row, column=0, sticky="w", pady=6)
            e = ttkb.Entry(f, width=34, font=FONT_BASE, bootstyle=ACCENT_STYLE, **kw)
            e.grid(row=row, column=1, padx=12, pady=6, sticky="w")
            setattr(self, attr, e)
        self._e_user.insert(0, _load_last_login())
        self._e_org.insert(0, str(ORG_ID_DEFAULT))
        self._e_pass.bind("<Return>", lambda _: self._do_login())
        self._e_user.bind("<Return>", lambda _: self._e_pass.focus_set())

        ttkb.Label(f, text="Hasło nie jest nigdzie zapisywane — tylko login.",
                   bootstyle="secondary", font=FONT_TINY).grid(row=3, column=0, columnspan=2, sticky="w", pady=(2, 12))
        bf = ttkb.Frame(f)
        bf.grid(row=4, column=0, columnspan=2, sticky="w")
        btn_login = ttkb.Button(bf, text="Zaloguj", command=self._do_login,
                                 bootstyle=ACCENT_STYLE, padding=(18, 7))
        btn_login.pack(side="left")
        self._lbl_auth = ttkb.Label(bf, text="", font=FONT_BASE)
        self._lbl_auth.pack(side="left", padx=12)

    # ── TAB 2: Import usług ───────────────────────────────────────────────────
    def _tab_import(self, p):
        outer = ttkb.Frame(p, padding=(16, 14))
        outer.pack(fill="both", expand=True)

        sec_file = ttkb.Labelframe(outer, text=" Plik Excel ", bootstyle=ACCENT_STYLE, padding=10)
        sec_file.pack(fill="x", pady=(0, 8))
        fr = ttkb.Frame(sec_file); fr.pack(fill="x")
        ttkb.Button(fr, text="📂 Wybierz plik...", command=self._choose_xlsx,
                    bootstyle="secondary", padding=(10, 5)).pack(side="left")
        self._lbl_xlsx = ttkb.Label(fr, text="— nie wybrano —", font=FONT_SMALL, bootstyle="secondary")
        self._lbl_xlsx.pack(side="left", padx=10)

        fr2 = ttkb.Frame(sec_file); fr2.pack(fill="x", pady=(8, 0))
        ttkb.Label(fr2, text="Arkusz:", font=FONT_SMALL).pack(side="left")
        self._sheet_var = tk.StringVar(value=DEFAULT_SHEET)
        self._cb_sheet = ttkb.Combobox(fr2, textvariable=self._sheet_var, state="readonly",
                                       width=20, font=FONT_SMALL, bootstyle=ACCENT_STYLE)
        self._cb_sheet.pack(side="left", padx=(6, 0))

        sec_params = ttkb.Labelframe(outer, text=" Parametry usług ", bootstyle=ACCENT_STYLE, padding=10)
        sec_params.pack(fill="x", pady=(0, 8))
        pr = ttkb.Frame(sec_params); pr.pack(fill="x")
        ttkb.Label(pr, text="Typ usługi (wszystkie):", font=FONT_SMALL).pack(side="left")
        self._e_typ = ttkb.Entry(pr, width=18, font=FONT_SMALL)
        self._e_typ.insert(0, DEFAULT_TYP_USLUGI)
        self._e_typ.pack(side="left", padx=(4, 14))
        ttkb.Label(pr, text="Pochodzenie usługi (wszystkie):", font=FONT_SMALL).pack(side="left")
        self._e_pochodzenie = ttkb.Entry(pr, width=14, font=FONT_SMALL)
        self._e_pochodzenie.insert(0, DEFAULT_POCHODZENIE)
        self._e_pochodzenie.pack(side="left", padx=4)
        ttkb.Label(sec_params,
                 text="Kategoria usługi jest brana osobno dla każdej usługi — z kolumny w Excelu.",
                 bootstyle="secondary", font=FONT_TINY).pack(anchor="w", pady=(6, 0))

        sec_go = ttkb.Labelframe(outer, text=" Podgląd i wysyłka ", bootstyle=ACCENT_STYLE, padding=10)
        sec_go.pack(fill="both", expand=True)
        gob = ttkb.Frame(sec_go); gob.pack(fill="x")
        ttkb.Button(gob, text="① 🔍 Pokaż podgląd", command=self._show_preview,
                    bootstyle="secondary", padding=(14, 6)).pack(side="left")
        send_style = "danger" if IS_PROD else "success"
        self._btn_send = ttkb.Button(gob, text="②  ▶  Wyślij usługi", command=self._do_send,
                  bootstyle=send_style, padding=(14, 6), state="disabled")
        self._btn_send.pack(side="left", padx=8)
        self._progress = ttkb.Progressbar(gob, length=240, mode="determinate", bootstyle=f"{ACCENT_STYLE}-striped")
        self._progress.pack(side="left", padx=12)
        self._lbl_progress = ttkb.Label(gob, text="", font=FONT_SMALL)
        self._lbl_progress.pack(side="left")
        ttkb.Label(sec_go, text="Kolejność: najpierw \"Pokaż podgląd\" (nic nie wysyła) — dopiero wtedy odblokuje się \"Wyślij usługi\".",
                 bootstyle="secondary", font=FONT_TINY).pack(anchor="w", pady=(4, 0))

        self._preview_text = ttkb.ScrolledText(
            sec_go, height=12, font=("Consolas", 9), bg="#f7f7f7", fg="#222",
            state="disabled", wrap="word", auto_hide=True)
        self._preview_text.pack(fill="both", expand=True, pady=(8, 0))

    # ── HELPERS ───────────────────────────────────────────────────────────────
    def _log_msg(self, msg, tag="inf"):
        from datetime import datetime as _dt
        ts = _dt.now().strftime("%H:%M:%S")
        self._log.configure(state="normal")
        self._log.insert("end", f"[{ts}] {msg}\n", tag)
        self._log.see("end")
        self._log.configure(state="disabled")
        logging.info(msg)

    def _bg(self, fn): return threading.Thread(target=fn, daemon=True).start()

    def _set_preview_text(self, text):
        self._preview_text.configure(state="normal")
        self._preview_text.delete("1.0", "end")
        self._preview_text.insert("end", text)
        self._preview_text.configure(state="disabled")

    # ── LOGOWANIE ─────────────────────────────────────────────────────────────
    def _do_login(self):
        user, pw = self._e_user.get().strip(), self._e_pass.get()
        try:
            org_id = int(self._e_org.get().strip())
        except ValueError:
            messagebox.showwarning("Logowanie", "Organizacja musi być liczbą (ID)."); return
        if not user or not pw:
            messagebox.showwarning("Logowanie", "Podaj login i hasło"); return
        self._lbl_auth.config(text="Logowanie...", bootstyle="secondary")

        def _t():
            ok = self.client.login(user, pw)
            if not ok:
                self.after(0, lambda: self._lbl_auth.config(text="Błąd logowania", bootstyle="danger"))
                self.after(0, lambda: self._log_msg("Błąd logowania — sprawdź login/hasło", "err"))
                return
            _save_last_login(user)
            if not self.client.select_organization(org_id):
                self.after(0, lambda: self._lbl_auth.config(
                    text=f"Zalogowano, ale nie udało się przełączyć na org {org_id}", bootstyle="danger"))
                self.after(0, lambda: self._log_msg(f"Nie udało się przełączyć na organizationId={org_id}", "err"))
                return
            self.after(0, lambda: self._lbl_auth.config(text="✓ Zalogowano", bootstyle="success"))
            self.after(0, lambda: self._lbl_status.config(text=f"●  Zalogowano (org {org_id})"))
            self.after(0, lambda: self._log_msg(f"Zalogowano jako {user}, organizacja {org_id}", "ok"))
            self.after(0, lambda: self._nb.select(1))
        self._bg(_t)

    # ── PLIK EXCEL ────────────────────────────────────────────────────────────
    def _choose_xlsx(self):
        path = filedialog.askopenfilename(
            title="Wybierz plik Excel", filetypes=[("Excel", "*.xlsx *.xlsm"), ("Wszystkie pliki", "*.*")])
        if not path:
            return
        self._xlsx_path = path
        self._lbl_xlsx.config(text=os.path.basename(path), bootstyle="success")
        try:
            sheets = list_sheet_names(path)
        except Exception as e:
            messagebox.showerror("Excel", f"Nie udało się otworzyć pliku:\n{e}")
            self._log_msg(f"Błąd otwierania {path}: {e}", "err")
            return
        self._cb_sheet.configure(values=sheets)
        if DEFAULT_SHEET in sheets:
            self._sheet_var.set(DEFAULT_SHEET)
        elif sheets:
            self._sheet_var.set(sheets[0])
        self._log_msg(f"Wybrano plik: {path} (arkusze: {', '.join(sheets)})", "inf")

    # ── PODGLĄD ───────────────────────────────────────────────────────────────
    def _show_preview(self):
        if not self._xlsx_path:
            messagebox.showwarning("Podgląd", "Najpierw wybierz plik Excel."); return
        sheet = self._sheet_var.get().strip()
        if not sheet:
            messagebox.showwarning("Podgląd", "Wybierz arkusz."); return
        typ_uslugi = self._e_typ.get().strip()
        pochodzenie = self._e_pochodzenie.get().strip()
        if not typ_uslugi or not pochodzenie:
            messagebox.showwarning("Podgląd", "Uzupełnij Typ usługi i Pochodzenie usługi."); return

        self._btn_send.config(state="disabled")
        self._set_preview_text("Wczytywanie Excela i słowników...")

        def _t():
            services, err = read_employee_services_from_excel(self._xlsx_path, sheet)
            if err:
                self.after(0, lambda: messagebox.showerror("Excel", err))
                self.after(0, lambda: self._set_preview_text(f"Błąd: {err}"))
                self.after(0, lambda: self._log_msg(f"Błąd odczytu Excela: {err}", "err"))
                return
            if not services:
                self.after(0, lambda: self._set_preview_text("Brak usług do wprowadzenia w tym arkuszu."))
                return

            structure = fetch_structure_elements(self.client, USLUGI_DICTIONARY_ID)
            if not structure:
                self.after(0, lambda: self._set_preview_text(
                    f"Błąd: nie udało się pobrać struktury słownika {USLUGI_DICTIONARY_ID}."))
                return
            self._structure_elements = structure

            existing = fetch_dictionary_values(self.client, USLUGI_DICTIONARY_ID)
            existing_contents = {(v.get("content") or "").strip().lower() for v in existing}

            to_create, skipped, errors = [], [], []
            dict_value_cache = {}
            for content, kategoria, row in services:
                if content.strip().lower() in existing_contents:
                    skipped.append(content)
                    continue
                _, err2 = resolve_value_attributes(
                    self.client, structure, kategoria, typ_uslugi, pochodzenie, dict_value_cache)
                if err2:
                    errors.append(f"{content}: {err2}")
                    continue
                to_create.append((content, kategoria, row))

            self.after(0, lambda: self._render_preview(to_create, skipped, errors, len(existing)))
        self._bg(_t)

    def _render_preview(self, to_create, skipped, errors, next_display_order):
        self._pending_services = to_create
        self._skipped_existing = skipped
        self._next_display_order = next_display_order
        lines = []
        lines.append(f"Typ usługi (wszystkie): {self._e_typ.get().strip()}")
        lines.append(f"Pochodzenie usługi (wszystkie): {self._e_pochodzenie.get().strip()}")
        lines.append("─" * 70)
        for content, kategoria, row in to_create:
            lines.append(f"  [wiersz {row}] {content}  (kategoria: {kategoria})")
        if skipped:
            lines.append("")
            lines.append(f"⚠ Pomijam (już istnieją w słowniku): {', '.join(skipped)}")
        if errors:
            lines.append("")
            lines.append("✗ Błędy (te usługi NIE zostaną utworzone):")
            for e in errors:
                lines.append(f"   {e}")
        lines.append("─" * 70)
        lines.append(f"RAZEM: {len(to_create)} nowych usług do utworzenia")
        self._set_preview_text("\n".join(lines))
        self._log_msg(f"Podgląd: {len(to_create)} nowych, {len(skipped)} pominiętych (już istnieją), "
                      f"{len(errors)} błędów", "inf")
        self._btn_send.config(state="normal" if to_create else "disabled")

    # ── WYSYŁKA ───────────────────────────────────────────────────────────────
    def _do_send(self):
        if not self._pending_services:
            messagebox.showwarning("Wysyłka", "Najpierw kliknij 'Pokaż podgląd'.")
            return
        n = len(self._pending_services)
        warn_prefix = "⚠ PRODUKCJA — usługi trafią od razu do systemu.\n\n" if IS_PROD else ""
        if not messagebox.askyesno("Potwierdzenie wysyłki",
            f"{warn_prefix}Dodać {n} usług do systemu?\nTej operacji nie można cofnąć hurtowo."):
            return

        services = list(self._pending_services)
        typ_uslugi = self._e_typ.get().strip()
        pochodzenie = self._e_pochodzenie.get().strip()
        structure = self._structure_elements
        display_order = self._next_display_order
        self._btn_send.config(state="disabled")
        self._progress.configure(maximum=len(services), value=0)
        self._lbl_progress.config(text="Wysyłanie...")

        def _t():
            ok = fail = 0
            dict_value_cache = {}
            for i, (content, kategoria, row) in enumerate(services, 1):
                attrs, err = resolve_value_attributes(
                    self.client, structure, kategoria, typ_uslugi, pochodzenie, dict_value_cache)
                if err:
                    fail += 1
                    self.after(0, lambda content=content, err=err: self._log_msg(f"✗ {content} — {err}", "err"))
                else:
                    success, err2 = post_service(self.client, content, display_order + i - 1, attrs)
                    if success:
                        ok += 1
                        self.after(0, lambda content=content: self._log_msg(f"✓ {content} — utworzono", "ok"))
                    else:
                        fail += 1
                        self.after(0, lambda content=content, err2=err2: self._log_msg(f"✗ {content} — {err2}", "err"))
                self.after(0, lambda i=i, ok=ok, fail=fail: (
                    self._progress.configure(value=i),
                    self._lbl_progress.config(text=f"{i}/{len(services)}  ✓ {ok}  ✗ {fail}")))
            summary = f"Zakończono: {ok}/{len(services)} utworzono, {fail} błędów"
            self.after(0, lambda: self._log_msg(f"— {summary} —", "ok" if fail == 0 else "err"))
            self.after(0, lambda: messagebox.showinfo("Gotowe", summary))
            self._pending_services = []
        self._bg(_t)


# ─────────────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    try:
        App().mainloop()
    except Exception as e:
        logging.error(f"Nieobsłużony wyjątek przy starcie okna: {e}\n{traceback.format_exc()}")
        _fatal_startup_error("Błąd uruchomienia okna", str(e), exc=e)
