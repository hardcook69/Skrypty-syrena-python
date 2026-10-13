#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Zestawienie POWIĄZAŃ zadań per DPS — z Excela importowego (Zdania grupowe …)
robi nowy plik Excel, w którym KAŻDY ARKUSZ to inny DPS, a w środku widać,
które zadanie pracownika jest połączone z którym zadaniem mieszkańca, w
rozbiciu na usługi.

Źródło: arkusz z kolumnami (dopasowanie po nagłówku):
  "Nazwa organizacji", "Usługa pracownik", "Zadanie pracownik",
  "Zadanie mieszkaniec", "Usługa mieszkaniec".

Układ każdego arkusza DPS (grupowane po usłudze pracownika):
  wiersz tytułu (nazwa DPS) → nagłówki kolumn →
  pasek "Usługa: <nazwa>" → pod nim pary [Zadanie pracownik | Zadanie
  mieszkaniec | Usługa mieszkaniec]. Zadanie pracownika bez odpowiednika
  mieszkańca: w kolumnie "Zadanie mieszkaniec" wstawiamy "—".

Tylko odczyt źródła. Użycie:
  python zestawienie_powiazan_per_dps.py --xlsx "Zdania grupowe … uzupelnione.xlsx"
  python zestawienie_powiazan_per_dps.py --xlsx plik.xlsx --sheet Arkusz1 --out wynik.xlsx
"""
import argparse
import re
import sys
from collections import defaultdict, OrderedDict

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("pip install openpyxl"); sys.exit(1)


# ── odczyt źródła (obsługa scalonych komórek) ─────────────────────────────────
def _build_merge_map(ws):
    m = {}
    for rng in ws.merged_cells.ranges:
        top_val = ws.cell(row=rng.min_row, column=rng.min_col).value
        for row in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                m[(row, col)] = top_val
    return m


def _cell_value(ws, merge_map, row, col):
    v = ws.cell(row=row, column=col).value
    if v is None:
        v = merge_map.get((row, col))
    return v


def _find_col(ws, header_row, target):
    for col in range(1, ws.max_column + 1):
        h = ws.cell(row=header_row, column=col).value
        if h and str(h).strip() == target:
            return col
    raise SystemExit(f"✗ Nie znaleziono kolumny '{target}' w wierszu nagłówka arkusza '{ws.title}'.")


def _dps_from_usluga(usluga):
    if not usluga:
        return None
    m = re.search(r'DPS\s*"([^"]+)"', str(usluga))
    return f'DPS "{m.group(1)}"' if m else None


def _clean(v):
    return str(v).strip() if v is not None and str(v).strip() else None


def read_rows(xlsx_path, sheet_name):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise SystemExit(f"✗ Arkusz '{sheet_name}' nie istnieje w {xlsx_path}. Dostępne: {wb.sheetnames}")
    ws = wb[sheet_name]
    mm = _build_merge_map(ws)
    c_org = _find_col(ws, 1, "Nazwa organizacji")
    c_up = _find_col(ws, 1, "Usługa pracownik")
    c_zp = _find_col(ws, 1, "Zadanie pracownik")
    c_zm = _find_col(ws, 1, "Zadanie mieszkaniec")
    c_um = _find_col(ws, 1, "Usługa mieszkaniec")

    rows = []
    for r in range(2, ws.max_row + 1):
        zp = _clean(_cell_value(ws, mm, r, c_zp))
        zm = _clean(_cell_value(ws, mm, r, c_zm))
        if not zp and not zm:
            continue
        up = _clean(_cell_value(ws, mm, r, c_up))
        rows.append({
            "dps": _clean(_cell_value(ws, mm, r, c_org)) or _dps_from_usluga(up) or "(bez DPS)",
            "usluga_pracownik": up or "(bez usługi)",
            "zadanie_pracownik": zp,
            "zadanie_mieszkaniec": zm,
            "usluga_mieszkaniec": _clean(_cell_value(ws, mm, r, c_um)),
        })
    return rows


# ── budowa wyniku ─────────────────────────────────────────────────────────────
_SHEETNAME_BAD = re.compile(r'[\[\]:*?/\\]')


def _sheet_name(dps, used):
    # "DPS \"św. Brata Alberta\"" -> "św. Brata Alberta"
    m = re.search(r'DPS\s*"([^"]+)"', dps or "")
    base = m.group(1) if m else (dps or "DPS")
    base = _SHEETNAME_BAD.sub(" ", base).strip()[:31] or "DPS"
    name, i = base, 2
    while name.lower() in used:
        suffix = f" ({i})"
        name = base[:31 - len(suffix)] + suffix
        i += 1
    used.add(name.lower())
    return name


def build_workbook(rows, out_path):
    # grupowanie: dps -> usługa pracownika -> lista par (dedup)
    by_dps = OrderedDict()
    for row in rows:
        by_dps.setdefault(row["dps"], defaultdict(list))
    # kolejność DPS: wg liczby wierszy malejąco, potem alfabetycznie
    counts = defaultdict(int)
    for row in rows:
        counts[row["dps"]] += 1
    dps_order = sorted(by_dps.keys(), key=lambda d: (-counts[d], d.lower()))

    seen = defaultdict(set)   # dps -> set krotek (dedup)
    grouped = {d: defaultdict(list) for d in by_dps}
    for row in rows:
        key = (row["usluga_pracownik"], row["zadanie_pracownik"], row["zadanie_mieszkaniec"])
        if key in seen[row["dps"]]:
            continue
        seen[row["dps"]].add(key)
        grouped[row["dps"]][row["usluga_pracownik"]].append(row)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    title_font = Font(bold=True, size=13, color="FFFFFF")
    title_fill = PatternFill("solid", fgColor="2F5597")
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="4472C4")
    grp_font = Font(bold=True, color="1F3864")
    grp_fill = PatternFill("solid", fgColor="D6E0F0")
    thin = Side(style="thin", color="B4C6E7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap_top = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="center")

    HEADERS = ["Zadanie pracownik", "Zadanie mieszkaniec", "Usługa mieszkaniec"]
    used_names = set()
    per_sheet = []

    for dps in dps_order:
        ws = wb.create_sheet(title=_sheet_name(dps, used_names))
        ws.column_dimensions["A"].width = 52
        ws.column_dimensions["B"].width = 52
        ws.column_dimensions["C"].width = 46

        # wiersz 1: tytuł DPS
        ws.merge_cells("A1:C1")
        tc = ws.cell(row=1, column=1, value=dps)
        tc.font = title_font; tc.fill = title_fill
        tc.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 22

        # wiersz 2: nagłówki kolumn
        for j, h in enumerate(HEADERS, start=1):
            c = ws.cell(row=2, column=j, value=h)
            c.font = hdr_font; c.fill = hdr_fill; c.alignment = center; c.border = border
        ws.freeze_panes = "A3"

        r = 3
        pair_count = 0
        for usluga in sorted(grouped[dps].keys(), key=str.lower):
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
            gc = ws.cell(row=r, column=1, value=f"Usługa: {usluga}")
            gc.font = grp_font; gc.fill = grp_fill
            gc.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            r += 1
            for row in sorted(grouped[dps][usluga],
                              key=lambda x: (x["zadanie_pracownik"] or "").lower()):
                ws.cell(row=r, column=1, value=row["zadanie_pracownik"] or "—").alignment = wrap_top
                ws.cell(row=r, column=2, value=row["zadanie_mieszkaniec"] or "—").alignment = wrap_top
                ws.cell(row=r, column=3, value=row["usluga_mieszkaniec"] or "").alignment = wrap_top
                for j in (1, 2, 3):
                    ws.cell(row=r, column=j).border = border
                r += 1
                pair_count += 1
            r += 1  # pusty wiersz-separator
        per_sheet.append((ws.title, pair_count))

    wb.save(out_path)
    return per_sheet


def main():
    ap = argparse.ArgumentParser(description="Zestawienie powiązań zadań per DPS (1 arkusz = 1 DPS).")
    ap.add_argument("--xlsx", required=True, help="Źródłowy Excel (Zdania grupowe … uzupelnione.xlsx).")
    ap.add_argument("--sheet", default="Arkusz1", help="Nazwa arkusza źródłowego (domyślnie 'Arkusz1').")
    ap.add_argument("--out", default="Powiazania_zadan_per_DPS.xlsx", help="Plik wynikowy.")
    args = ap.parse_args()

    rows = read_rows(args.xlsx, args.sheet)
    if not rows:
        print("Brak wierszy z zadaniami w źródle — nic do zrobienia."); sys.exit(0)
    per_sheet = build_workbook(rows, args.out)

    total = sum(n for _, n in per_sheet)
    print(f"✓ Zapisano: {args.out}")
    print(f"  Arkuszy (DPS): {len(per_sheet)}  |  łącznie powiązań: {total}")
    for name, n in per_sheet:
        print(f"   - {name}: {n}")


if __name__ == "__main__":
    main()
