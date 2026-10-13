#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Arkusz DECYZYJNY korekty „opiekun medyczny".

Z głównego Excela importowego (Zdania grupowe … uzupelnione.xlsx) znajduje
zadania, w których CZYNNOŚĆ OPIEKUŃCZA (kąpiel, toaleta, posiłki, karmienie,
dopajanie, pościel, naczynia, strzyżenie/paznokcie) jest błędnie przypisana do
roli personelu MEDYCZNEGO (<OPIEKUN MEDYCZNY> / <PIELĘGNIARKA>), i buduje
arkusz do przejrzenia pozycja po pozycji, z dwiema opcjami korekty obok siebie:
  • Opcja A — tylko ROLA (ta sama usługa, rola → OPIEKUN),
  • Opcja B — ROLA + USŁUGA (usługa dobrana wg treści zadania).
Użytkownik wpisuje ewentualne wyjątki w pustych kolumnach i odsyła.

Tylko ODCZYT źródła. Użycie:
  python przygotuj_korekte_opiekun_medyczny.py --xlsx "…uzupelnione.xlsx"
  python przygotuj_korekte_opiekun_medyczny.py --xlsx plik.xlsx --sheet Arkusz1 --out decyzje.xlsx
"""
import argparse
import re
import sys

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    print("pip install openpyxl"); sys.exit(1)


# ── odczyt źródła (scalone komórki) ───────────────────────────────────────────
def _build_merge_map(ws):
    m = {}
    for rng in ws.merged_cells.ranges:
        top = ws.cell(row=rng.min_row, column=rng.min_col).value
        for row in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                m[(row, col)] = top
    return m


def _cell(ws, mm, r, c):
    v = ws.cell(row=r, column=c).value
    if v is None:
        v = mm.get((r, c))
    return str(v).strip() if v is not None and str(v).strip() else ""


def _find_col(ws, target):
    # tolerancja na końcowy dwukropek / białe znaki (np. nagłówek 'Stanowisko:')
    t = target.strip().rstrip(":").strip().casefold()
    for col in range(1, ws.max_column + 1):
        h = ws.cell(row=1, column=col).value
        if h and str(h).strip().rstrip(":").strip().casefold() == t:
            return col
    raise SystemExit(f"✗ Nie znaleziono kolumny '{target}' w nagłówku arkusza '{ws.title}'.")


# ── reguły ────────────────────────────────────────────────────────────────────
ROLE_RE = re.compile(r"<([^>]+)>")
MED_ROLES = {"OPIEKUN MEDYCZNY", "PIELĘGNIARKA"}

# słowa-klucze czynności opiekuńczych (treść zadania pracownika)
CARE_KW = ["kąpiel", "toalet", "ubran", "rozebran", "posiłek", "posilek",
           "śniadan", "sniadan", "obiad", "kolacj", "karmienie", "dopajanie",
           "naczyń", "naczyn", "pościel", "posciel", "strzyżenie", "strzyzenie",
           "paznokc", "golenie"]

# treść zadania -> docelowa baza usługi (Opcja B)
USLUGA_B_RULES = [
    (("kąpiel",), "kąpiel"),
    (("toalet", "ubran", "rozebran"), "Toalety"),
    (("posiłek", "posilek", "śniadan", "sniadan", "obiad", "kolacj",
      "karmienie", "dopajanie", "naczyń", "naczyn"), "Wyżywienie"),
    (("pościel", "posciel"), "Sprzątanie"),
    (("strzyżenie", "strzyzenie", "paznokc", "golenie"), "Toalety"),
]


def _role_of(usluga):
    m = ROLE_RE.search(usluga)
    return m.group(1).strip() if m else None


def _is_care(zadanie):
    z = zadanie.lower()
    return any(k in z for k in CARE_KW)


def _usluga_base_B(zadanie):
    z = zadanie.lower()
    for keys, base in USLUGA_B_RULES:
        if any(k in z for k in keys):
            return base
    return None


def _rewrite_usluga(usluga, new_base=None, new_role="OPIEKUN"):
    """Zamienia rolę w <...> na new_role; jeśli new_base podany, zamienia też
    część przed '<' (nazwę usługi), zachowując '(…) - DPS "X"' po '>'."""
    m = re.match(r"^(.*?)<([^>]+)>(.*)$", usluga)
    if not m:
        return usluga
    base, _role, tail = m.group(1), m.group(2), m.group(3)
    if new_base is not None:
        base = new_base + " "
    return f"{base}<{new_role}>{tail}"


# ── główna logika ─────────────────────────────────────────────────────────────
def collect(xlsx_path, sheet_name):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise SystemExit(f"✗ Arkusz '{sheet_name}' nie istnieje. Dostępne: {wb.sheetnames}")
    ws = wb[sheet_name]
    mm = _build_merge_map(ws)
    c_org = _find_col(ws, "Nazwa organizacji")
    c_up = _find_col(ws, "Usługa pracownik")
    c_zp = _find_col(ws, "Zadanie pracownik")
    c_st = _find_col(ws, "Stanowisko")

    seen = set()
    items = []
    for r in range(2, ws.max_row + 1):
        up = _cell(ws, mm, r, c_up)
        zp = _cell(ws, mm, r, c_zp)
        if not up or not zp:
            continue
        role = _role_of(up)
        if role not in MED_ROLES:
            continue
        if not _is_care(zp):
            continue
        key = (_cell(ws, mm, r, c_org), up, zp)
        if key in seen:
            continue
        seen.add(key)
        base_b = _usluga_base_B(zp)
        items.append({
            "dps": _cell(ws, mm, r, c_org),
            "zadanie": zp,
            "usluga": up,
            "stanowisko": _cell(ws, mm, r, c_st),
            "opcja_a": _rewrite_usluga(up, new_base=None, new_role="OPIEKUN"),
            "opcja_b": _rewrite_usluga(up, new_base=base_b, new_role="OPIEKUN"),
        })
    items.sort(key=lambda x: (x["dps"], x["usluga"].lower(), x["zadanie"].lower()))
    return items


def build(items, out_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Korekta opiekun medyczny"

    title_font = Font(bold=True, size=12, color="FFFFFF")
    title_fill = PatternFill("solid", fgColor="2F5597")
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="4472C4")
    fill_a = PatternFill("solid", fgColor="E2EFDA")   # Opcja A - zielonkawy
    fill_b = PatternFill("solid", fgColor="FFF2CC")   # Opcja B - żółtawy
    fill_in = PatternFill("solid", fgColor="FCE4D6")  # do wypełnienia - łososiowy
    thin = Side(style="thin", color="B4C6E7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers = [
        ("Lp.", 5), ("DPS", 16), ("Zadanie pracownik", 50),
        ("Obecna usługa (z rolą)", 40), ("Obecne Stanowisko", 30),
        ("Proponowana rola", 14),
        ("OPCJA A — tylko rola", 40), ("OPCJA B — rola + usługa", 40),
        ("Twoja rola (jeśli inna)", 16), ("Wybór usługi (A/B/własna)", 18),
        ("Uwagi", 24),
    ]

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    tc = ws.cell(row=1, column=1,
                 value=('Korekta roli „opiekun medyczny” dla czynności opiekuńczych. '
                        'Przejrzyj pozycje; wypełnij kolumny łososiowe tylko dla WYJĄTKÓW '
                        '(domyślnie: rola = OPIEKUN; wybór usługi A albo B).'))
    tc.font = title_font; tc.fill = title_fill
    tc.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30

    for j, (h, w) in enumerate(headers, start=1):
        c = ws.cell(row=2, column=j, value=h)
        c.font = hdr_font; c.fill = hdr_fill; c.alignment = center; c.border = border
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    ws.freeze_panes = "A3"

    r = 3
    for i, it in enumerate(items, start=1):
        vals = [i, it["dps"], it["zadanie"], it["usluga"], it["stanowisko"],
                "OPIEKUN", it["opcja_a"], it["opcja_b"], "", "", ""]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.alignment = center if j in (1, 6) else wrap
            c.border = border
            if j == 7:
                c.fill = fill_a
            elif j == 8:
                c.fill = fill_b
            elif j in (9, 10, 11):
                c.fill = fill_in
        r += 1

    wb.save(out_path)
    return len(items)


def main():
    ap = argparse.ArgumentParser(description="Arkusz decyzyjny korekty 'opiekun medyczny'.")
    ap.add_argument("--xlsx", required=True, help="Źródłowy Excel (Zdania grupowe … uzupelnione.xlsx).")
    ap.add_argument("--sheet", default="Arkusz1", help="Arkusz źródłowy (domyślnie 'Arkusz1').")
    ap.add_argument("--out", default="Korekta_opiekun_medyczny_DECYZJE.xlsx", help="Plik wynikowy.")
    args = ap.parse_args()

    items = collect(args.xlsx, args.sheet)
    if not items:
        print("Nie znaleziono pozycji do korekty."); sys.exit(0)
    n = build(items, args.out)
    print(f"✓ Zapisano: {args.out}  ({n} pozycji do decyzji)")
    diff = sum(1 for it in items if it["opcja_a"] != it["opcja_b"])
    print(f"  Pozycji, gdzie Opcja A ≠ Opcja B (zmiana także usługi): {diff}")
    for it in items:
        mark = "  ≠" if it["opcja_a"] != it["opcja_b"] else "   "
        print(f"{mark} [{it['dps'][4:]}] {it['zadanie'][:55]}")


if __name__ == "__main__":
    main()
