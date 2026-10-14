#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Raport niezgodności „USŁUGA ≠ treść zadania" z propozycją właściwej usługi.

Z głównego Excela wylicza dla każdego zadania USŁUGĘ wynikającą z TREŚCI
(klasyfikator słów-kluczy) i porównuje z usługą obecnie przypisaną. Gdy się
różnią (i treść pewnie wskazuje inną, konkretną usługę) — wpisuje pozycję do
arkusza decyzyjnego z proponowaną usługą. Wysoka precyzja: gdy treść nie
pasuje pewnie do żadnej konkretnej usługi, NIE flagujemy (brak fałszywych
alarmów). Rola/stanowisko bez zmian.

Tylko odczyt. Użycie:
  python raport_niezgodnosci_usluga.py --xlsx "Zdania_grupowe_POPRAWIONY.xlsx"
  python raport_niezgodnosci_usluga.py --xlsx plik.xlsx --sheet Arkusz1 --out raport.xlsx
"""
import argparse
import re
import sys

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    print("pip install openpyxl"); sys.exit(1)

import przygotuj_korekte_opiekun_medyczny as K  # _build_merge_map/_cell/_find_col/_role_of/_rewrite_usluga

# Reguły treść→USŁUGA (kolejność ma znaczenie: pierwsze trafienie wygrywa).
USLUGA_RULES = [
    (("strzyżenie", "strzyzenie", "paznokc", "golenie"), "Toalety"),
    (("kąpiel", "kapiel"), "kąpiel"),
    # UWAGA: słowa dot. UBIERANIA muszą być wąskie, żeby NIE łapać
    # „Układanie ubrań/odzieży" (to porządki = Sprzątanie, nie Toalety).
    (("toalet", "pieluchomajtk", "pielucha", "pieluch",
      "ubraniu", "ubieran", "rozebran", "rozbieran"), "Toalety"),
    (("posiłk", "posilk", "śniadan", "sniadan", "obiad", "kolacj", "podwieczor",
      "karmienie", "dopajanie", "naczyń", "naczyn", "kanap", "stołów", "stolow"), "Wyżywienie"),
    (("pranie", "prania", "mycie okien", "ręcznik", "recznik", "pościel", "posciel",
      "sprzątan", "sprzatan", "porządk", "porzadk", "szafek", "dezynfekc", "czyszczen"), "Sprzątanie"),
    (("okresowa ocena", "ocena stanu zdrowia"), "Ocena stanu zdrowia"),
    (("monitorowanie stanu", "pomiar glikemii", "pomiar ciśnien", "pomiar cisnien",
      "saturacj", "parametr"), "Monitorowanie"),
    (("opatrun", "zabiegi pielęgnacyjne", "zabiegi pielegnacyjne"), "Zabiegi pielęgnacyjne"),
    (("podawanie lek", "podanie lek", "rozkładanie lek", "rozkladanie lek",
      "spis braku lek", "zażyciu lek", "zazyciu lek", "recept"), "Leki"),
    (("doprowadz", "odprowadz", "sprowadz", "do wyjazdu", "do przychodni",
      "wyjazd z mieszka"), "Doprowadzanie"),
]


def proposed_base(zadanie):
    z = (zadanie or "").lower()
    for keys, base in USLUGA_RULES:
        if any(k in z for k in keys):
            return base
    return None


def _usluga_base(usluga):
    m = re.match(r"^(.*?)\s*<", usluga or "")
    base = (m.group(1).strip() if m else (usluga or "").split(" - DPS")[0].strip())
    return base


def _norm_base(b):
    # 'Sprzątanie dom' i 'Sprzątanie' = ta sama rodzina
    bl = (b or "").strip().lower()
    if bl.startswith("sprzątanie") or bl.startswith("sprzatanie"):
        return "sprzątanie"
    return bl


def collect(xlsx_path, sheet):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if sheet not in wb.sheetnames:
        raise SystemExit(f"✗ Brak arkusza '{sheet}'. Dostępne: {wb.sheetnames}")
    ws = wb[sheet]
    mm = K._build_merge_map(ws)
    c_org = K._find_col(ws, "Nazwa organizacji")
    c_up = K._find_col(ws, "Usługa pracownik")
    c_zp = K._find_col(ws, "Zadanie pracownik")

    seen = set()
    items = []
    for r in range(2, ws.max_row + 1):
        up = K._cell(ws, mm, r, c_up)
        zp = K._cell(ws, mm, r, c_zp)
        if not up or not zp:
            continue
        prop = proposed_base(zp)
        if prop is None:
            continue
        cur_base = _usluga_base(up)
        if _norm_base(cur_base) == _norm_base(prop):
            continue   # usługa już pasuje do treści — nie flagujemy
        role = K._role_of(up)
        prop_full = K._rewrite_usluga(up, new_base=prop, new_role=role) if role else None
        note = "" if role else "obecna usługa nie ma roli <…> — uzupełnij ręcznie"
        key = (K._cell(ws, mm, r, c_org), up, zp)
        if key in seen:
            continue
        seen.add(key)
        items.append({
            "dps": K._cell(ws, mm, r, c_org), "zp": zp, "up": up,
            "role": role or "(brak)",
            "prop": prop_full or (prop + " <…>"),
            "note": note,
        })
    items.sort(key=lambda x: (x["dps"], _usluga_base(x["up"]).lower(), x["zp"].lower()))
    return items


def build(items, out_path):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Niezgodności usługa"
    tf = Font(bold=True, size=12, color="FFFFFF"); tfill = PatternFill("solid", fgColor="2F5597")
    hf = Font(bold=True, color="FFFFFF"); hfill = PatternFill("solid", fgColor="4472C4")
    prop_fill = PatternFill("solid", fgColor="E2EFDA")
    in_fill = PatternFill("solid", fgColor="FCE4D6")
    thin = Side(style="thin", color="B4C6E7"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top"); center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    heads = [("Lp.", 5), ("DPS", 16), ("Zadanie pracownik", 52), ("Obecna usługa", 40),
             ("Rola", 16), ("PROPONOWANA usługa", 40), ("Twoja usługa (jeśli inna)", 26), ("Uwagi", 30)]
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(heads))
    tc = ws.cell(row=1, column=1,
                 value=("Niezgodności „usługa ≠ treść zadania”. Przejrzyj; jeśli proponowana usługa "
                        "jest OK — nic nie rób. Dla wyjątków wpisz własną w kolumnie łososiowej."))
    tc.font = tf; tc.fill = tfill; tc.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28
    for j, (h, w) in enumerate(heads, start=1):
        c = ws.cell(row=2, column=j, value=h); c.font = hf; c.fill = hfill; c.alignment = center; c.border = border
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    ws.freeze_panes = "A3"
    r = 3
    for i, it in enumerate(items, start=1):
        vals = [i, it["dps"], it["zp"], it["up"], it["role"], it["prop"], "", it["note"]]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.alignment = center if j in (1, 5) else wrap
            c.border = border
            if j == 6:
                c.fill = prop_fill
            elif j == 7:
                c.fill = in_fill
        r += 1
    wb.save(out_path)
    return len(items)


def main():
    ap = argparse.ArgumentParser(description="Raport niezgodności usługa↔treść z propozycją.")
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--sheet", default="Arkusz1")
    ap.add_argument("--out", default="Niezgodnosci_usluga_DECYZJE.xlsx")
    args = ap.parse_args()
    items = collect(args.xlsx, args.sheet)
    n = build(items, args.out)
    print(f"✓ Zapisano: {args.out}  ({n} niezgodności usługa↔treść)")
    from collections import Counter
    c = Counter((_usluga_base(it["up"]), it["prop"].split(" <")[0]) for it in items)
    print("  Rozkład (obecna baza → proponowana : ile):")
    for (o, p), k in sorted(c.items(), key=lambda x: -x[1]):
        print(f"    {k:3d}  {o}  →  {p}")


if __name__ == "__main__":
    main()
